# -*- coding: utf-8 -*-
"""
Build the AI Medical Assist final documentation in the PUCIT Final Documentation Format.
Output: d:\\fyp\\Rohan_Hassan\\AI_Medical_Assist_Documentation.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BUILD = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BUILD)
OUT = os.path.join(ROOT, "AI_Medical_Assist_Documentation.docx")

FONT = "Times New Roman"
INK = RGBColor(0x1B, 0x27, 0x33)
ACCENT = RGBColor(0x1F, 0x3B, 0x5C)
GREY = RGBColor(0x54, 0x65, 0x7A)

FIG = {n: os.path.join(BUILD, f) for n, f in {
    "arch": "fig1_architecture.png",
    "pipeline": "fig2_ml_pipeline.png",
    "multitask": "fig3_multitask.png",
    "apparch": "fig4_app_arch.png",
    "gantt": "fig5_gantt.png",
}.items()}

PROJECT_TITLE = "AI Medical Assist"
PROJECT_SUBTITLE = ("A Predictive Lifestyle Health-Risk Assessment System for "
                    "Diabetes and Chronic Kidney Disease")
HEADER_SHORT = "AI Medical Assist  -  Predictive Health-Risk Assessment"
FOOTER_TEXT = "Punjab University College of Information Technology"

# ----------------------------------------------------------------------- helpers

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_field(paragraph, instr):
    """Insert a Word field (e.g. PAGE, TOC) into a paragraph."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = ""
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(b); r.append(it); r.append(sep); r.append(t); r.append(e)
    return run


def style_runs(paragraph, size=12, bold=False, italic=False, color=INK):
    for r in paragraph.runs:
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        rpr = r._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(a), FONT)


def para_border_top(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "4"); top.set(qn("w:color"), "808080")
    pbdr.append(top); pPr.append(pbdr)


def para_border_bottom(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "808080")
    pbdr.append(bottom); pPr.append(pbdr)


def set_pgnum(section, fmt, start=1):
    sectPr = section._sectPr
    old = sectPr.find(qn("w:pgNumType"))
    if old is not None:
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)      # "lowerRoman" or "decimal"
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


# ----------------------------------------------------------------------- doc setup
doc = Document()

# base Normal style
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.font.color.rgb = INK
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
pf.space_after = Pt(8)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# heading styles per PUCIT spec
def conf_heading(name, size, italic=False):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.italic = italic
    st.font.color.rgb = ACCENT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:cs"), FONT)
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

conf_heading("Heading 1", 14)
conf_heading("Heading 2", 12)
conf_heading("Heading 3", 12, italic=True)

# caption style
try:
    cap = doc.styles["Caption"]
except KeyError:
    cap = doc.styles.add_style("Caption", 1)
cap.font.name = FONT
cap.font.size = Pt(10)
cap.font.italic = True
cap.font.bold = False
cap.font.color.rgb = GREY
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(4)
cap.paragraph_format.space_after = Pt(10)

# A4 + margins on first (title) section
sec0 = doc.sections[0]
sec0.page_height = Inches(11.69)
sec0.page_width = Inches(8.27)
sec0.top_margin = Inches(1.0)
sec0.bottom_margin = Inches(1.0)
sec0.left_margin = Inches(1.25)
sec0.right_margin = Inches(1.0)


# ----------------------------------------------------------------------- content helpers

def body(text, align="justify", size=12, bold=False, italic=False, space_after=8,
         color=INK):
    p = doc.add_paragraph()
    p.add_run(text)
    amap = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    p.alignment = amap[align]
    p.paragraph_format.space_after = Pt(space_after)
    style_runs(p, size=size, bold=bold, italic=italic, color=color)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r1 = p.add_run(bold_lead);
        r2 = p.add_run(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in (r1, r2):
            r.font.name = FONT; r.font.size = Pt(12); r.font.color.rgb = INK
            rpr = r._element.get_or_add_rPr(); rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), FONT)
            rpr.append(rf)
        r1.font.bold = True
    else:
        p.add_run(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style_runs(p, size=12)
    p.paragraph_format.space_after = Pt(4)
    return p


def heading(text, level):
    h = doc.add_heading(text, level=level)
    return h


def front_title(text):
    p = doc.add_paragraph()
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    style_runs(p, size=16, bold=True, color=ACCENT)
    return p


_fig_count = 0
_tab_count = 0

def add_figure(img_key, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(FIG[img_key], width=Inches(width))
    # caption with SEQ field => "Figure X: caption"
    cap_p = doc.add_paragraph(style="Caption")
    r = cap_p.add_run("Figure ")
    # SEQ Figure
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = " SEQ Figure \\* ARABIC "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    rr = cap_p.add_run()._r
    rr.append(b); rr.append(it); rr.append(sep); rr.append(t); rr.append(e)
    cap_p.add_run(":  " + caption)
    for rn in cap_p.runs:
        rn.font.name = FONT; rn.font.size = Pt(10); rn.font.italic = True
        rn.font.color.rgb = GREY
    return cap_p


def add_table_caption(caption):
    cap_p = doc.add_paragraph(style="Caption")
    cap_p.add_run("Table ")
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = " SEQ Table \\* ARABIC "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    rr = cap_p.add_run()._r
    rr.append(b); rr.append(it); rr.append(sep); rr.append(t); rr.append(e)
    cap_p.add_run(":  " + caption)
    for rn in cap_p.runs:
        rn.font.name = FONT; rn.font.size = Pt(10); rn.font.italic = True
        rn.font.color.rgb = GREY


def make_table(headers, rows, widths=None, caption=None):
    if caption:
        add_table_caption(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.add_run(h)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_runs(p, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_bg(hdr[i], "1F3B5C")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.add_run(str(val))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_runs(p, size=11)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


# ======================================================================= TITLE PAGE
def spacer(n=1):
    for _ in range(n):
        doc.add_paragraph()

body(FOOTER_TEXT, align="center", size=14, bold=True, color=ACCENT, space_after=2)
body("University of the Punjab, Lahore", align="center", size=12, color=GREY, space_after=2)
spacer(2)
body("Final Year Design Project", align="center", size=13, italic=True, color=GREY, space_after=2)
spacer(1)
body(PROJECT_TITLE, align="center", size=24, bold=True, color=ACCENT, space_after=4)
body(PROJECT_SUBTITLE, align="center", size=13, italic=True, color=INK, space_after=6)
spacer(2)
body("Session:  BSIT  (2022 - 2026)", align="center", size=12, bold=True, space_after=8)
spacer(1)
body("Project Advisor", align="center", size=12, bold=True, color=ACCENT, space_after=2)
body("Prof. Hafiz Shahzad", align="center", size=12, space_after=6)
body("Head of Department", align="center", size=12, bold=True, color=ACCENT, space_after=2)
body("Prof. Muhammad Fahim", align="center", size=12, space_after=8)
body("Submitted By", align="center", size=12, bold=True, color=ACCENT, space_after=2)
body("Muhammad Hassan          Roll No. 2716", align="center", size=12, space_after=2)
body("Muhammad Rohan Ashraf     Roll No. 2720", align="center", size=12, space_after=8)
spacer(2)
body("Punjab University College of Information Technology", align="center", size=12, bold=True, space_after=1)
body("University of the Punjab, Lahore.", align="center", size=12, space_after=1)


# ======================================================================= FRONT MATTER SECTION (roman)
doc.add_section(WD_SECTION.NEW_PAGE)
sec_front = doc.sections[-1]
sec_front.page_height = Inches(11.69); sec_front.page_width = Inches(8.27)
sec_front.top_margin = Inches(1.0); sec_front.bottom_margin = Inches(1.0)
sec_front.left_margin = Inches(1.25); sec_front.right_margin = Inches(1.0)
sec_front.header.is_linked_to_previous = False
sec_front.footer.is_linked_to_previous = False
set_pgnum(sec_front, "lowerRoman", start=1)

# header
hp = sec_front.header.paragraphs[0]
hp.text = HEADER_SHORT
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_runs(hp, size=9, color=GREY)
para_border_bottom(hp)
# footer with text + page number
fp = sec_front.footer.paragraphs[0]
fp.text = ""
r = fp.add_run(FOOTER_TEXT)
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_runs(fp, size=9, color=GREY)
para_border_top(fp)
# tab + page field on the right
tabp = sec_front.footer.add_paragraph()
tabp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_field(tabp, " PAGE ")
style_runs(tabp, size=9, color=GREY)

# ---- Statement of Submission
front_title("Statement of Submission")
body("This is to certify that Mr. Muhammad Hassan, Roll No. 2716, and Mr. Muhammad Rohan "
     "Ashraf, Roll No. 2720, have successfully completed the final year design project titled "
     "“AI Medical Assist – A Predictive Lifestyle Health-Risk Assessment System for "
     "Diabetes and Chronic Kidney Disease” at the Punjab University College of Information "
     "Technology, University of the Punjab, Lahore, in the partial fulfilment of the requirements "
     "for the degree of Bachelor of Science in Information Technology (BSIT).")
spacer(3)
p = doc.add_paragraph(); p.add_run("______________________________"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_runs(p, size=12)
body("Project Advisor", size=12, bold=True, space_after=0)
body("Prof. Hafiz Shahzad", size=12, space_after=0)
body("Punjab University College of Information Technology, Lahore", size=12, color=GREY)
spacer(2)
p = doc.add_paragraph(); p.add_run("______________________________"); style_runs(p, size=12)
body("Head of Department", size=12, bold=True, space_after=0)
body("Prof. Muhammad Fahim", size=12, space_after=0)
body("Punjab University College of Information Technology, Lahore", size=12, color=GREY)

# ---- Proofreading Certificate
doc.add_page_break()
front_title("Proofreading Certificate")
body("It is to certify that I have read this document meticulously and circumspectly. I am "
     "convinced that the resultant document does not contain any spelling, punctuation or "
     "grammatical mistakes. Overall, I find this document well organised, and I am in no doubt "
     "that its objectives have been successfully met.")
spacer(4)
p = doc.add_paragraph(); p.add_run("______________________________"); style_runs(p, size=12)
body("Business Communication and Technical Writing", size=12, bold=True, space_after=0)
body("Lecturer, PUCIT", size=12, color=GREY)

# ---- Acknowledgement
doc.add_page_break()
front_title("Acknowledgement")
body("All praise and gratitude are due to Almighty Allah, whose blessings enabled us to "
     "undertake and complete this project. We owe our sincere appreciation to our supervisors, "
     "Prof. Hafiz Shahzad and Prof. Muhammad Fahim, whose guidance, patience and constructive "
     "feedback shaped this work at every stage. Their insistence that we question our own "
     "assumptions – particularly around the design of the machine learning core – pushed "
     "the project towards a genuinely meaningful and defensible solution.")
body("We are equally thankful to the faculty of the Punjab University College of Information "
     "Technology for providing the academic foundation that made this work possible. Finally, we "
     "thank our families and friends, whose quiet and unwavering support carried us through the "
     "demanding months of this project.")
spacer(2)
body("Muhammad Hassan", size=12, bold=True, space_after=0)
body("Muhammad Rohan Ashraf", size=12, bold=True)

# ---- Abstract
doc.add_page_break()
front_title("Abstract")
body("Diabetes Mellitus and Chronic Kidney Disease are among the most widespread and costly "
     "chronic illnesses in the world, yet most existing health applications only react after a "
     "person is already ill – typically by reading a laboratory blood report. By that point, "
     "the opportunity for prevention has often passed. AI Medical Assist takes the opposite "
     "approach. Instead of confirming disease from clinical lab values, it predicts a person’s "
     "future risk of developing diabetes and kidney disease from their everyday lifestyle – "
     "factors such as physical activity, diet, body-mass index, sleep, smoking, blood-pressure "
     "history and family history – without requiring any blood test.")
body("At the heart of the system is a multi-task machine learning model that predicts both "
     "diseases at once. Because diabetes is a leading cause of kidney disease, a single model with "
     "a shared representation learns the physiological link between the two conditions, rather than "
     "treating them as unrelated problems. The model is trained on the CDC Behavioural Risk Factor "
     "Surveillance System (BRFSS) dataset, which contains the lifestyle and health records of "
     "roughly 250,000 real individuals. To ensure the predictions are trustworthy, the system "
     "explains every result using SHAP feature attribution and goes one step further with a "
     "“what-if” counterfactual engine that tells the user exactly which habits to change "
     "to lower their risk.")
body("The platform is delivered as a mobile-first application built with Flutter and served by a "
     "Django REST Framework backend, with a Gemini conversational interface for plain-language "
     "health guidance and Google Maps integration for locating nearby medical facilities. The "
     "system is positioned strictly as a preventive screening and educational tool, not a "
     "diagnostic authority.")
body("Keywords:  Preventive Healthcare, Lifestyle Risk Prediction, Multi-Task Learning, "
     "Explainable AI, Counterfactual Recommendations.", italic=True, color=GREY)

# ---- Table of Contents
doc.add_page_break()
front_title("Table of Contents")
p = doc.add_paragraph()
add_field(p, ' TOC \\o "1-3" \\h \\z \\u ')
style_runs(p, size=12)

# ---- List of Figures
doc.add_page_break()
front_title("List of Figures")
p = doc.add_paragraph()
add_field(p, ' TOC \\h \\z \\c "Figure" ')
style_runs(p, size=12)

# ---- List of Tables
doc.add_page_break()
front_title("List of Tables")
p = doc.add_paragraph()
add_field(p, ' TOC \\h \\z \\c "Table" ')
style_runs(p, size=12)


# ======================================================================= BODY SECTION (arabic)
doc.add_section(WD_SECTION.NEW_PAGE)
sec_body = doc.sections[-1]
sec_body.page_height = Inches(11.69); sec_body.page_width = Inches(8.27)
sec_body.top_margin = Inches(1.0); sec_body.bottom_margin = Inches(1.0)
sec_body.left_margin = Inches(1.25); sec_body.right_margin = Inches(1.0)
sec_body.header.is_linked_to_previous = False
sec_body.footer.is_linked_to_previous = False
set_pgnum(sec_body, "decimal", start=1)
hp = sec_body.header.paragraphs[0]
hp.text = HEADER_SHORT; hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_runs(hp, size=9, color=GREY); para_border_bottom(hp)
fp = sec_body.footer.paragraphs[0]
fp.add_run(FOOTER_TEXT); fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_runs(fp, size=9, color=GREY); para_border_top(fp)
tabp = sec_body.footer.add_paragraph(); tabp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_field(tabp, " PAGE "); style_runs(tabp, size=9, color=GREY)

# ----------------------------------------------------------------- 1. Project Title
heading("1   Project Title", 1)
body("AI Medical Assist – A Predictive Lifestyle Health-Risk Assessment System for Diabetes "
     "and Chronic Kidney Disease.")
body("The project delivers an intelligent, mobile-first health companion that estimates a "
     "person’s future risk of developing Diabetes Mellitus (DM) and Chronic Kidney Disease "
     "(CKD) from their daily lifestyle, explains the reasoning behind every prediction, and "
     "recommends concrete, personalised actions to reduce that risk.")

# ----------------------------------------------------------------- 2. Project Overview Statement
heading("2   Project Overview Statement", 1)

heading("2.1   The Problem", 2)
body("Chronic diseases such as diabetes and kidney disease develop silently over many years. In "
     "developing regions in particular, people rarely undergo regular laboratory screening, so the "
     "condition is usually discovered only once symptoms or complications appear. The vast majority "
     "of existing health applications reinforce this reactive pattern: they require a blood test or "
     "clinical lab value and then simply confirm what an already-present report indicates. This is "
     "of limited value, because by the time a lab result is abnormal, the disease is frequently "
     "already established. The genuine, unmet need is for a tool that warns people earlier – "
     "before the disease takes hold – using information they already have about their own "
     "lifestyle.")

heading("2.2   The Proposed Solution", 2)
body("AI Medical Assist reframes the goal from detection to prevention. A user answers a short, "
     "approachable questionnaire about their everyday routine – activity level, diet, "
     "body-mass index, sleep, smoking and alcohol habits, blood-pressure history and family "
     "history. From these lifestyle signals alone, and with no blood test required, the system "
     "predicts the user’s future risk of both diabetes and kidney disease, presents the result "
     "as an easy-to-understand risk tier, explains which factors are driving that risk, and "
     "recommends specific changes that would measurably lower it.")
body("The platform is deliberately designed as a two-tier engine. The first tier performs "
     "lifestyle-only screening that anyone can use immediately. The optional second tier allows a "
     "user who already has laboratory values to enter them and receive a sharper, refined estimate. "
     "This keeps the product widely accessible while still rewarding clinical data when it is "
     "available.")

heading("2.3   The Machine Learning Innovation", 2)
body("The intelligence of the system is a multi-task learning model that predicts both diseases "
     "simultaneously. Because diabetes is one of the leading causes of kidney disease, the two "
     "conditions are physiologically connected. Rather than building two isolated models, AI "
     "Medical Assist uses a single model with a shared representation that learns this connection "
     "directly from the data, with two dedicated output “heads” – one for diabetes "
     "risk and one for kidney risk. This design is both medically faithful and technically "
     "stronger, as knowledge learned for one disease improves the prediction of the other.")
body("Crucially, the model is trained to predict outcomes that cannot be trivially derived from a "
     "fixed formula. It learns genuine, non-linear risk patterns from the CDC Behavioural Risk "
     "Factor Surveillance System (BRFSS) dataset – a national health survey covering "
     "approximately 250,000 real individuals – rather than re-encoding a clinical threshold "
     "rule. Because serious disease is comparatively rare in the population, the data is "
     "deliberately rebalanced, and the model is judged on recall (sensitivity), AUC-ROC and F1 "
     "score rather than on raw accuracy, which would otherwise hide missed high-risk cases.")
body("Finally, the system does not stop at a risk number. Using SHAP it identifies the factors "
     "most responsible for each individual’s risk, and a counterfactual “what-if” "
     "engine translates those factors into actionable advice – for example, indicating how "
     "much a user’s risk could fall if they increased weekly physical activity or reduced "
     "their body-mass index. This turns a passive warning into a practical, motivating plan.")

heading("2.4   Project Objectives", 2)
make_table(
    ["ID", "Objective", "Success Measure"],
    [
        ["OBJ-1", "Build a multi-task model predicting future diabetes and kidney risk from lifestyle data", "Strong recall and AUC-ROC on held-out data"],
        ["OBJ-2", "Train and validate on the large, real-world BRFSS dataset with class balancing", "Robust performance on rare positive cases"],
        ["OBJ-3", "Explain every prediction with SHAP feature attribution", "Top contributing factors shown per result"],
        ["OBJ-4", "Provide what-if counterfactual recommendations", "Actionable risk-reduction guidance generated"],
        ["OBJ-5", "Deliver a secure Flutter mobile app with JWT authentication", "Working login, questionnaire and dashboard"],
        ["OBJ-6", "Integrate Gemini chat and Google Maps doctor lookup", "Plain-language guidance and nearby facilities"],
    ],
    widths=[0.8, 3.4, 2.2],
    caption="Project objectives and their success measures.",
)

# ----------------------------------------------------------------- 3. High Level System Components
heading("3   High Level System Components", 1)
body("AI Medical Assist is organised as a layered, modular system in which each component has a "
     "single clear responsibility. The presentation layer collects input and displays results; the "
     "application layer orchestrates business logic and security; the intelligence layer performs "
     "prediction, explanation and advice; and the data and external-service layers provide "
     "persistence and third-party capabilities. Figure 1 presents the high-level architecture.")
add_figure("arch", "High-level layered architecture of AI Medical Assist.", width=6.4)

heading("3.1   Mobile Application (Flutter)", 2)
body("The user-facing application is built with Flutter using a clean, layered structure that "
     "separates presentation, domain and data concerns. It guides the user through the lifestyle "
     "questionnaire, presents risk tiers and explanations in an uncluttered dashboard, and hosts "
     "the conversational chat experience. The interface emphasises ample whitespace, consistent "
     "typography and thumb-friendly navigation suitable for everyday, non-specialist users.")

heading("3.2   Backend and API Layer (Django REST Framework)", 2)
body("The backend is a modular monolith built on the Django REST Framework. It exposes secure "
     "REST endpoints protected by JWT authentication and is internally divided into well-defined "
     "services – user and authentication, prediction, explainability and advisory – each "
     "following the single-responsibility principle. This layered organisation keeps the system "
     "maintainable and easy to defend, while avoiding the operational overhead of a premature "
     "microservices design.")

heading("3.3   Predictive ML Engine", 2)
body("The predictive engine is the core of the system. It loads the trained multi-task model once "
     "at start-up and serves real-time risk predictions for both diabetes and kidney disease. "
     "Figure 2 shows the end-to-end machine learning pipeline, from lifestyle input through to "
     "explained, actionable output, and Figure 3 illustrates the multi-task model architecture in "
     "which a shared backbone feeds two disease-specific risk heads.")
add_figure("pipeline", "Predictive machine learning pipeline (lifestyle input to actionable advice).", width=6.6)
add_figure("multitask", "Multi-task model: a shared backbone with two disease-specific risk heads.", width=6.0)

heading("3.4   Explainability and Advisory", 2)
body("Trust is essential in any health-facing system, so no prediction is presented as an "
     "unexplained black box. The explainability service applies SHAP to reveal the factors most "
     "responsible for each user’s risk, while the advisory service uses a counterfactual "
     "what-if engine to recommend the specific lifestyle changes that would most reduce that risk. "
     "A Gemini-powered conversational interface then communicates these insights in clear, "
     "supportive, plain language.")

heading("3.5   Data and Geolocation Services", 2)
body("A PostgreSQL database securely stores user accounts, questionnaire history and prediction "
     "records, enabling users to track how their risk evolves over time. The Google Maps service "
     "helps users who are flagged as high-risk locate nearby clinics or doctors, closing the loop "
     "between awareness and action.")

# ----------------------------------------------------------------- 4. Optional functional units
heading("4   List of Optional Functional Units", 1)
body("The following units enhance the system but are not essential to its core preventive "
     "function. They are scoped as optional so that the central prediction-and-advice experience "
     "can be delivered and validated first.")
make_table(
    ["Optional Unit", "Description"],
    [
        ["Clinical (Tier-2) Refinement", "Lets users who have lab values enter them to sharpen the lifestyle-based risk estimate."],
        ["Risk History and Trends", "Visualises how a user’s predicted risk changes over repeated assessments."],
        ["PDF Health Report Export", "Generates a shareable summary of risk, contributing factors and recommendations."],
        ["Reminders and Notifications", "Encourages periodic reassessment and adherence to recommended habit changes."],
        ["Multi-Language Support", "Provides the questionnaire and guidance in local languages for wider accessibility."],
        ["Doctor Lookup via Maps", "Surfaces nearby medical facilities for users flagged as high-risk."],
    ],
    widths=[2.2, 4.2],
    caption="Optional functional units and their descriptions.",
)

# ----------------------------------------------------------------- 5. Application Architecture
heading("5   Application Architecture", 1)
body("The application follows a modular-monolith architecture with a clean, layered internal "
     "structure, applying SOLID and DRY principles throughout. The Flutter client communicates "
     "with the Django REST backend exclusively over secure HTTPS using JWT-authenticated "
     "requests. The backend routes each request to the appropriate service, invokes the ML engine "
     "for prediction and the explainability and advisory components for interpretation, and "
     "persists results to PostgreSQL. External capabilities – Gemini for conversation and "
     "Google Maps for geolocation – are integrated as cleanly isolated services. Figure 4 "
     "shows the request flow across these components.")
add_figure("apparch", "Application architecture and request flow across major components.", width=6.4)
body("This architecture was chosen deliberately. For a two-person team delivering within a single "
     "academic timeline, a modular monolith provides the clarity, testability and clean separation "
     "of a well-structured system without the deployment and coordination cost of microservices. "
     "The ML engine is isolated behind a clean interface so that the underlying model can be "
     "improved or replaced without disturbing the rest of the system. Table 3 summarises the "
     "technology stack.")
make_table(
    ["Layer / Concern", "Technology", "Rationale"],
    [
        ["Mobile Frontend", "Flutter, Dart", "Cross-platform, clean and responsive UI"],
        ["Backend and API", "Django REST Framework, JWT", "Secure, modular, maintainable API layer"],
        ["ML Engine", "scikit-learn / gradient boosting, multi-task model", "Lifestyle-based dual-disease risk prediction"],
        ["Explainability", "SHAP, counterfactual what-if engine", "Transparent, actionable results"],
        ["Database", "PostgreSQL", "Secure relational storage of users and history"],
        ["Conversational AI", "Gemini API", "Plain-language health guidance"],
        ["Geolocation", "Google Maps API", "Nearby doctor and facility lookup"],
    ],
    widths=[1.9, 2.3, 2.2],
    caption="Technology stack mapped to system concerns.",
)

# ----------------------------------------------------------------- 6. Gantt chart
heading("6   Gantt Chart", 1)
body("The project is planned across a sixteen-week timeline. Development is deliberately "
     "front-loaded onto the machine learning core – the component carrying the greatest "
     "technical risk and the greatest academic value – before backend, mobile and integration "
     "work proceed. Figure 5 presents the schedule.")
add_figure("gantt", "Project schedule across the sixteen-week timeline.", width=6.6)
make_table(
    ["Phase", "Weeks", "Milestone"],
    [
        ["Requirements & Dataset Preparation", "1 - 2", "Problem framing and dataset finalised"],
        ["ML Model Development (multi-task)", "3 - 7", "Trained and validated risk model"],
        ["Explainability & What-If Engine", "5 - 7", "SHAP and counterfactual advice working"],
        ["Backend Development (Django REST)", "6 - 8", "Secure REST APIs ready"],
        ["Mobile App Development (Flutter)", "8 - 10", "Questionnaire and dashboard integrated"],
        ["Gemini Chat & Maps Integration", "10 - 12", "Conversational guidance and doctor lookup"],
        ["Testing & Evaluation", "13 - 14", "All modules validated"],
        ["Documentation & Deployment", "15 - 16", "Project finalised and submitted"],
    ],
    widths=[2.8, 1.0, 2.6],
    caption="Project phases, timeline and milestones.",
)

# ----------------------------------------------------------------- update fields on open
settings = doc.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
settings.append(upd)

doc.save(OUT)
print("SAVED:", OUT)
